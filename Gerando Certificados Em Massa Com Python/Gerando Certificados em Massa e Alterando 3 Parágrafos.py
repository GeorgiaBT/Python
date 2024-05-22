from docx import Document
from docx.shared import Pt

from openpyxl import load_workbook
import os


arquivoalunos = "C:\\Users\\georg\\Documents\\Python - GitHub\\Python\\Gerando Certificados Em Massa Com Python\\DadosAlunos.xlsx"
planilhadadosaluno = load_workbook(arquivoalunos)

sheet_selecionada = planilhadadosaluno["Nomes"]

for linha in range(2, len(sheet_selecionada["A"]) + 1):

    arquivoword = Document("C:\\Users\\georg\\Documents\\Python - GitHub\\Python\\Gerando Certificados Em Massa Com Python\\Certificado3.docx")

    estilo = arquivoword.styles["Normal"]

    #pegando o nome do aluno qnd passar na celula
    nomealuno = sheet_selecionada['A%s' % linha].value
    dia = sheet_selecionada['B%s' % linha].value
    mes = sheet_selecionada['C%s' % linha].value
    ano = sheet_selecionada['D%s' % linha].value
    nomecurso = sheet_selecionada['E%s' % linha].value
    nomeinstrutor = sheet_selecionada['F%s' % linha].value

    for paragrafo in arquivoword.paragraphs:

        if "@nome" in paragrafo.text:
            paragrafo.text = nomealuno
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

        paragrafop1 = "Concluiu com sucesso o curso de"
        paragrafop2 = ", como carga horária de 20 horas, promovido pela escola de Cursos Online em "
        paragrafocompleto = f"{paragrafop1} {nomecurso} {paragrafop2} {dia} de {mes} de {ano}"

        if "escola" in paragrafo.text:
            paragrafo.text = paragrafocompleto
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

        if "Instrutor" in paragrafo.text:
            paragrafo.text = nomeinstrutor + " - Instrutor"
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

    caminhocertificados = "C:\\Users\\georg\\Downloads\\" + nomealuno + ".docx"
    arquivoword.save(caminhocertificados)